"""文档解析。实现 core.spi.DocumentParser。

🔴 ``parse`` 故意是**同步**的：诚实反映它 CPU 密集，逼调用方显式 to_thread。
🔴 认不出的格式直接拒绝，不做"尽力而为"的解码——
   错误的文本比没有文本更糟：它会被嵌入、被检索、被当成答案依据，且全程不报错。
"""
from __future__ import annotations

import io
from collections.abc import Callable
from pathlib import PurePath

from ...core.errors import UnsupportedDocument
from ...core.telemetry import get_logger
from ...core.types import ParsedDocument

logger = get_logger(__name__)


def _pdf(raw: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [(p.extract_text() or "") for p in reader.pages]
    return ParsedDocument(
        text="\n\n".join(p for p in pages if p.strip()),
        format="pdf", metadata={"pages": len(reader.pages)},
    )


def _docx(raw: bytes) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格常是知识库里最有价值的部分，不能漏
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return ParsedDocument(
        text="\n".join(blocks), format="docx",
        metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
    )


def _html(raw: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return ParsedDocument(
        text=soup.get_text("\n", strip=True), format="html",
        metadata={"title": (soup.title.string or "").strip() if soup.title else ""},
    )


def _text(raw: bytes) -> ParsedDocument:
    # 只有纯文本才允许解码。GB18030 覆盖中文 Windows 导出的常见编码。
    for enc in ("utf-8", "gb18030"):
        try:
            return ParsedDocument(text=raw.decode(enc), format="text", metadata={"encoding": enc})
        except UnicodeDecodeError:
            continue
    raise UnsupportedDocument("文本编码无法识别（已尝试 utf-8 与 gb18030）")


_BY_EXT: dict[str, Callable[[bytes], ParsedDocument]] = {
    ".pdf": _pdf, ".docx": _docx, ".html": _html, ".htm": _html,
    ".txt": _text, ".md": _text, ".markdown": _text, ".csv": _text, ".json": _text,
}
_MAGIC = [(b"%PDF-", _pdf, "pdf"), (b"PK\x03\x04", _docx, "docx")]


def supported_extensions() -> list[str]:
    return sorted(_BY_EXT)


class FileParser:
    def parse(self, raw: bytes, filename: str = "", content_type: str = "") -> ParsedDocument:
        if not raw:
            raise UnsupportedDocument("文件内容为空")

        fn = _BY_EXT.get(PurePath(filename).suffix.lower())
        if fn is None:
            for magic, parser, label in _MAGIC:
                if raw.startswith(magic):
                    logger.info("按魔数识别为 %s（文件名 %r 无有效扩展名）", label, filename)
                    fn = parser
                    break
        if fn is None and content_type:
            if "pdf" in content_type:
                fn = _pdf
            elif "wordprocessingml" in content_type:
                fn = _docx
            elif "html" in content_type:
                fn = _html
            elif content_type.startswith("text/"):
                fn = _text
        if fn is None:
            raise UnsupportedDocument(
                f"无法识别的文档格式（文件名 {filename!r}，Content-Type {content_type!r}）。"
                f"支持: {', '.join(supported_extensions())}",
                filename=filename,
            )

        try:
            doc = fn(raw)
        except UnsupportedDocument:
            raise
        except Exception as exc:
            # 不要把 pypdf / python-docx 的内部异常类型漏给上层
            raise UnsupportedDocument(
                f"文档解析失败（{type(exc).__name__}）：文件可能损坏或受密码保护",
                filename=filename,
            ) from exc

        if doc.is_empty:
            raise UnsupportedDocument(
                f"{doc.format} 解析结果为空（可能是扫描件，需要 OCR）",
                filename=filename, format=doc.format,
            )
        return doc
